"""
Generate friday_demo_slides.pptx for the BioImage hackathon BEH demo.

Three speakers, nine slides. Speaker badge (colored corner tag) on each
slide so the team can see at a glance whose slide is whose.

Run:
    python generate_slides.py

Output:
    /dcs/pg25/u1898019/Desktop/hackathon-beh-code/slides/friday_demo_slides.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# === Paths ===
IMG_DIR = Path('/dcs/pg25/u1898019/public_html/hackathon-beh')
OUT = Path('/dcs/pg25/u1898019/Desktop/hackathon-beh-code/slides/friday_demo_slides.pptx')
OUT.parent.mkdir(parents=True, exist_ok=True)

# === Colours ===
COL_WT = RGBColor(0x1f, 0x77, 0xb4)
COL_KO = RGBColor(0xff, 0x7f, 0x0e)
COL_KI = RGBColor(0xd6, 0x27, 0x28)
COL_BADEER = RGBColor(0x38, 0xa1, 0x69)
COL_EDWARD = RGBColor(0x31, 0x82, 0xce)
COL_HAILI  = RGBColor(0xd5, 0x3f, 0x8c)
COL_ALL    = RGBColor(0x4a, 0x55, 0x68)
COL_DARK   = RGBColor(0x2d, 0x37, 0x48)
COL_BODY   = RGBColor(0x2d, 0x37, 0x48)
COL_META   = RGBColor(0x71, 0x80, 0x96)
COL_BG_LIGHT = RGBColor(0xf7, 0xfa, 0xfc)
COL_QUOTE_BG = RGBColor(0xff, 0xfa, 0xf0)

# === Setup ===
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]  # blank layout


# === Helpers ===
def add_textbox(slide, left, top, width, height, text, *,
                font_size=18, bold=False, color=COL_BODY, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    return tb


def add_richtext(slide, left, top, width, height, paragraphs, *,
                 default_size=16, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of (text, {'bold': bool, 'color': RGBColor, 'size': pt, 'align': align, 'bullet': bool})
       where each item describes ONE paragraph; runs separated inside a paragraph by passing a list of (text, opts)."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.05)

    for i, item in enumerate(paragraphs):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        if isinstance(item, tuple) and len(item) == 2 and isinstance(item[0], str):
            text, opts = item
            opts = opts or {}
            p.alignment = opts.get('align', PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = text
            run.font.size = Pt(opts.get('size', default_size))
            run.font.bold = opts.get('bold', False)
            run.font.color.rgb = opts.get('color', COL_BODY)
            run.font.name = 'Calibri'
        elif isinstance(item, list):
            # Multi-run paragraph
            opts0 = item[0][1] if item else {}
            p.alignment = (opts0 or {}).get('align', PP_ALIGN.LEFT)
            for text, opts in item:
                opts = opts or {}
                run = p.add_run()
                run.text = text
                run.font.size = Pt(opts.get('size', default_size))
                run.font.bold = opts.get('bold', False)
                run.font.color.rgb = opts.get('color', COL_BODY)
                run.font.name = 'Calibri'
    return tb


def add_speaker_badge(slide, who):
    """who: 'badeer' | 'edward' | 'haili' | 'all' | 'cs' (Edward+Haili)"""
    if who == 'cs':
        # Two-tone badge: half edward, half haili
        x, y = SW - Inches(2.6), Inches(0.18)
        h = Inches(0.36)
        # Left half (Edward blue)
        sh1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(1.3), h)
        sh1.fill.solid(); sh1.fill.fore_color.rgb = COL_EDWARD; sh1.line.fill.background()
        # Right half (Haili pink)
        sh2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(1.3), y, Inches(1.3), h)
        sh2.fill.solid(); sh2.fill.fore_color.rgb = COL_HAILI; sh2.line.fill.background()
        # Text overlay
        tb = slide.shapes.add_textbox(x, y, Inches(2.6), h)
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = 'EDWARD + HAILI'
        run.font.size = Pt(11); run.font.bold = True
        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff); run.font.name = 'Calibri'
        return

    palette = {
        'badeer': (COL_BADEER, '🧬 BADEER'),
        'edward': (COL_EDWARD, '🛠 EDWARD'),
        'haili':  (COL_HAILI,  '📈 HAILI'),
        'all':    (COL_ALL,    'ALL THREE'),
    }
    col, label = palette[who]
    x, y = SW - Inches(2.0), Inches(0.18)
    w, h = Inches(1.7), Inches(0.36)
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = col; sh.line.fill.background()
    sh.adjustments[0] = 0.4
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(12); run.font.bold = True
    run.font.color.rgb = RGBColor(0xff, 0xff, 0xff); run.font.name = 'Calibri'


def add_title(slide, text, *, top=Inches(0.65)):
    return add_textbox(slide, Inches(0.5), top, SW - Inches(1), Inches(0.7),
                       text, font_size=28, bold=True, color=COL_DARK)


def add_image(slide, path, left, top, width=None, height=None):
    if width and height:
        return slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    if width:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    if height:
        return slide.shapes.add_picture(str(path), left, top, height=height)
    return slide.shapes.add_picture(str(path), left, top)


def add_filled_box(slide, left, top, width, height, fill_color, text=None,
                   font_size=14, font_color=COL_BODY, bold=False, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.adjustments[0] = 0.08
    sh.fill.solid(); sh.fill.fore_color.rgb = fill_color
    sh.line.color.rgb = RGBColor(0xe2, 0xe8, 0xf0)
    if text:
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]; p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size); run.font.bold = bold
        run.font.color.rgb = font_color; run.font.name = 'Calibri'
    return sh


def add_table(slide, left, top, width, height, data, *,
              header=True, font_size=12, col_widths=None):
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = w
    for r, row_data in enumerate(data):
        for c, val in enumerate(row_data):
            cell = tbl.cell(r, c)
            cell.text = ''
            tf = cell.text_frame
            tf.margin_left = tf.margin_right = Inches(0.07)
            tf.margin_top = tf.margin_bottom = Inches(0.04)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(font_size)
            run.font.name = 'Calibri'
            if header and r == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                cell.fill.solid()
                cell.fill.fore_color.rgb = COL_DARK
            else:
                run.font.color.rgb = COL_BODY
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COL_BG_LIGHT
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xff, 0xff, 0xff)
    return table_shape


# ===========================================================================
# Slide 1: Title (all three)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'all')

add_textbox(s, Inches(0.5), Inches(1.4), SW - Inches(1), Inches(1.0),
            'BioImage Hackathon BEH', font_size=46, bold=True, color=COL_DARK,
            align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.5), Inches(2.5), SW - Inches(1), Inches(0.8),
            'A point mutation that tries harder and fails more',
            font_size=28, bold=True, color=COL_KI, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.5), Inches(3.6), SW - Inches(1), Inches(0.6),
            'Cell migration analysis pipeline + cross-batch validation',
            font_size=20, color=COL_BODY, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.5), Inches(4.1), SW - Inches(1), Inches(0.6),
            'on 60 GFP-Arp2/3 time-lapse recordings',
            font_size=20, color=COL_BODY, align=PP_ALIGN.CENTER)

# Three name pills
name_y = Inches(5.3)
name_h = Inches(0.55)
gap = Inches(0.3)
pill_w = Inches(3.5)
total_w = pill_w * 3 + gap * 2
start_x = (SW - total_w) // 2
for i, (name, role, col) in enumerate([
    ('🧬  Badeer Ummat', 'biology', COL_BADEER),
    ('🛠  Edward Offord', 'CS', COL_EDWARD),
    ('📈  Haili Yuan', 'CS', COL_HAILI),
]):
    x = start_x + i * (pill_w + gap)
    add_filled_box(s, x, name_y, pill_w, name_h, col,
                   text=f'{name}  ·  {role}',
                   font_size=15, font_color=RGBColor(0xff, 0xff, 0xff), bold=True)

add_textbox(s, Inches(0.5), Inches(6.4), SW - Inches(1), Inches(0.4),
            'Warwick CAMDU  ·  2026-04-29 to 05-01',
            font_size=14, color=COL_META, align=PP_ALIGN.CENTER)


# ===========================================================================
# Slide 2: Biology + question (Badeer)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'badeer')
add_title(s, 'The biology: a gene that coordinates pushing and gripping')

# Left column
left_x, left_w = Inches(0.5), Inches(6.5)
add_textbox(s, left_x, Inches(1.5), left_w, Inches(0.4),
            'Cell migration is a four-step process:',
            font_size=18, bold=True, color=COL_DARK)

steps = [
    '1.  Push the front edge forward (lamellipodium).',
    '2.  Anchor the leading edge to the surface.',
    '3.  Generate pulling force.',
    '4.  Release the rear.',
]
for i, step in enumerate(steps):
    add_textbox(s, left_x + Inches(0.1), Inches(1.95) + Inches(0.42)*i, left_w - Inches(0.1), Inches(0.4),
                step, font_size=15, color=COL_BODY)

add_textbox(s, left_x, Inches(3.85), left_w, Inches(1.2),
            'The gene we perturb coordinates steps 1 and 2: it links protrusion to substrate adhesion. Without it, the cell can still try to push, but the push and the grip stop happening together.',
            font_size=14, color=COL_BODY)

add_textbox(s, left_x, Inches(5.1), left_w, Inches(1.2),
            'The marker:  GFP fused to Arp2/3, the protein complex that builds the lamellipodium. Brighter pixels at the cell edge = more push machinery being recruited.',
            font_size=14, color=COL_BODY)

# Right column
right_x, right_w = Inches(7.3), Inches(5.5)
add_textbox(s, right_x, Inches(1.5), right_w, Inches(0.4),
            'Three conditions, 30 cells:',
            font_size=18, bold=True, color=COL_DARK)

cond_items = [
    ('WT', COL_WT, 'wild type. The gene works.'),
    ('KO', COL_KO, 'knock-out. The gene is fully deleted.'),
    ('KI', COL_KI, 'knock-in. The gene is replaced by a single-amino-acid point-mutation version.'),
]
for i, (badge, col, text) in enumerate(cond_items):
    y = Inches(2.0) + Inches(0.65)*i
    # Badge box
    add_filled_box(s, right_x, y, Inches(0.7), Inches(0.4), col,
                   text=badge, font_size=14, font_color=RGBColor(0xff,0xff,0xff), bold=True)
    add_textbox(s, right_x + Inches(0.85), y - Inches(0.04), right_w - Inches(0.9), Inches(0.6),
                text, font_size=14, color=COL_BODY, anchor=MSO_ANCHOR.MIDDLE)

# Question quote
quote_y = Inches(4.05)
add_filled_box(s, right_x, quote_y, right_w, Inches(1.5), COL_QUOTE_BG)
add_textbox(s, right_x + Inches(0.15), quote_y + Inches(0.1), right_w - Inches(0.3), Inches(0.4),
            'The question:', font_size=14, bold=True, color=COL_KI)
add_textbox(s, right_x + Inches(0.15), quote_y + Inches(0.5), right_w - Inches(0.3), Inches(0.95),
            'Does the KI point mutation behave the same as a complete deletion (KO), or does swapping one amino acid produce a phenotype of its own?',
            font_size=14, color=COL_BODY)

add_textbox(s, right_x, Inches(5.7), right_w, Inches(1.0),
            '2D epifluorescence  ·  0.318 µm/px  ·  60 s/frame  ·  single channel  ·  16-bit  ·  two imaging sessions per condition',
            font_size=11, color=COL_META)


# ===========================================================================
# Slide 3: Pipeline + training (Edward)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'edward')
add_title(s, 'Pipeline journey: 75× reduction in segmentation failures')

# Pipeline table
pipeline_data = [
    ['Stage', 'What changed', 'Empty mask rate'],
    ['v1', 'Cellpose cyto, raw images', '9.7%'],
    ['v2', 'cyto3 + 1-99 percentile normalize on every frame', '7.9% (high-batch regression)'],
    ['v3', 'Adaptive normalize: per-stack p99 detector decides whether to stretch', '2.83%'],
    ['fine-tuned', 'Self-trained cyto3 on 586 filtered v3 outputs (no hand-drawn data)', '0.13%'],
]
add_table(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(2.3),
          pipeline_data, font_size=13,
          col_widths=[Inches(1.6), Inches(7.7), Inches(3.0)])

# Highlight final row's empty rate in KI red
# (python-pptx makes per-cell colour fiddly to retro-fit; the table style is enough.)

# Two-column lower section
add_textbox(s, Inches(0.5), Inches(4.05), Inches(7.5), Inches(0.4),
            'How the self-training works:',
            font_size=17, bold=True, color=COL_DARK)
steps_st = [
    '1.  Run v3 on all frames.',
    '2.  Filter: keep frames where mask area ≥ 30% of per-cell median AND IoU ≥ 0.5 with previous frame.',
    '3.  Drop the 2 worst files (ko1/4 95% empty, ko1/8 50% empty).',
    '4.  Treat the surviving 586 (frame, mask) pairs as pseudo-ground-truth.',
    '5.  Fine-tune cyto3 for 100 epochs on those pairs.',
]
for i, step in enumerate(steps_st):
    add_textbox(s, Inches(0.6), Inches(4.55) + Inches(0.42)*i, Inches(7.5), Inches(0.4),
                step, font_size=13, color=COL_BODY)

# Right side: stat boxes
stat_x = Inches(8.5)
stat_w = Inches(4.3)
stat_h = Inches(0.85)
stat_gap = Inches(0.15)
for i, (num, label) in enumerate([
    ('12 min', 'fine-tune on RTX A5000'),
    ('10 min', 're-run inference on 30 cells'),
    ('26 MB', 'final weights file (shipped on GitHub)'),
]):
    y = Inches(4.05) + i * (stat_h + stat_gap)
    add_filled_box(s, stat_x, y, stat_w, stat_h, COL_BG_LIGHT)
    add_textbox(s, stat_x, y + Inches(0.05), stat_w, Inches(0.45),
                num, font_size=24, bold=True, color=COL_DARK, align=PP_ALIGN.CENTER)
    add_textbox(s, stat_x, y + Inches(0.5), stat_w, Inches(0.35),
                label, font_size=12, color=COL_META, align=PP_ALIGN.CENTER)


# ===========================================================================
# Slide 4: Migration trajectories (Haili)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'haili')
add_title(s, 'Finding 1: KI cells migrate 35% slower')

# Left side: table + commentary
mig_data = [
    ['Condition', 'mean speed (µm/min)'],
    ['WT (N=10)', '1.03 ± 0.33'],
    ['KO (N=10)', '1.06 ± 0.29'],
    ['KI (N=10)', '0.68 ± 0.08'],
]
add_table(s, Inches(0.5), Inches(1.5), Inches(5.5), Inches(1.7), mig_data,
          font_size=13, col_widths=[Inches(2.2), Inches(3.3)])

add_textbox(s, Inches(0.5), Inches(3.35), Inches(5.8), Inches(0.45),
            "Cohen's d:  KI vs WT −1.40,  KI vs KO −1.68",
            font_size=14, bold=True, color=COL_KI)
add_textbox(s, Inches(0.5), Inches(3.8), Inches(5.8), Inches(0.45),
            '(both very large effects)',
            font_size=12, color=COL_META)

add_textbox(s, Inches(0.5), Inches(4.35), Inches(5.8), Inches(1.0),
            'KO migrates at WT speed. Only KI slows down. The point mutation breaks something the complete deletion does not.',
            font_size=14, color=COL_BODY)

add_textbox(s, Inches(0.5), Inches(5.45), Inches(5.8), Inches(0.4),
            'Two filters that matter:',
            font_size=14, bold=True, color=COL_DARK)
add_textbox(s, Inches(0.5), Inches(5.85), Inches(5.8), Inches(1.5),
            '• Mask-stability filter rejects per-pair frames with IoU < 0.3 or sudden area drops. Caught a fragmentation artefact in wt1/33 that had inflated speed by 5×.\n\n• Persistence capped at 1.0 (geometric upper bound).',
            font_size=12, color=COL_BODY)

# Right side: trajectory image
img_path = IMG_DIR / 'migration_trajectories_finetuned.png'
if img_path.exists():
    add_image(s, img_path, Inches(6.7), Inches(1.5), height=Inches(5.0))
add_textbox(s, Inches(6.7), Inches(6.6), Inches(6.4), Inches(0.5),
            'Per-cell trajectories, 30 cells across 6 folders. KI panels (right column) cover noticeably less ground.',
            font_size=10, color=COL_META, align=PP_ALIGN.CENTER)


# ===========================================================================
# Slide 5: Lamellipodia + dominant-negative (Haili)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'haili')
add_title(s, 'Finding 2: KI over-polarises Arp2/3 at the leading edge')

# Left side
lam_data = [
    ['Condition', 'lam / cyto intensity ratio'],
    ['WT', '1.53 ± 0.29'],
    ['KO', '1.47 ± 0.45'],
    ['KI', '1.83 ± 0.33'],
]
add_table(s, Inches(0.5), Inches(1.5), Inches(5.7), Inches(1.7), lam_data,
          font_size=13, col_widths=[Inches(1.7), Inches(4.0)])

add_textbox(s, Inches(0.5), Inches(3.35), Inches(5.9), Inches(0.45),
            "Cohen's d:  KI vs WT +0.90,  KI vs KO +0.85",
            font_size=14, bold=True, color=COL_KI)
add_textbox(s, Inches(0.5), Inches(3.8), Inches(5.9), Inches(0.45),
            '(both large effects)',
            font_size=12, color=COL_META)

add_textbox(s, Inches(0.5), Inches(4.35), Inches(5.9), Inches(2.0),
            'Per cell, brighter pixels inside the mask (above the per-frame median) are treated as lamellipodia. KI cells push more Arp2/3 to the edge than WT or KO.',
            font_size=14, color=COL_BODY)

# Right side: dominant-negative reading
right_x = Inches(6.7)
right_w = Inches(6.2)

quote_y = Inches(1.5)
add_filled_box(s, right_x, quote_y, right_w, Inches(2.0), COL_QUOTE_BG)
add_textbox(s, right_x + Inches(0.2), quote_y + Inches(0.15), right_w - Inches(0.4), Inches(0.4),
            'Combined biology reading:',
            font_size=15, bold=True, color=COL_KI)
add_textbox(s, right_x + Inches(0.2), quote_y + Inches(0.55), right_w - Inches(0.4), Inches(1.4),
            'KI cells push the front edge harder (more Arp2/3 polarisation, finding 2) yet move 35% less (finding 1). Effort goes in, locomotion does not come out.',
            font_size=14, color=COL_BODY)

add_textbox(s, right_x, Inches(3.7), right_w, Inches(1.5),
            'This is what "broken protrusion-adhesion coordination" predicts: the lamellipodium activates, but the cell cannot anchor or generate traction, so the push wastes energy.',
            font_size=14, color=COL_BODY)

add_textbox(s, right_x, Inches(5.4), right_w, Inches(1.5),
            'Dominant-negative reading: a broken-coordination mutation is functionally MORE costly than completely deleting the gene.',
            font_size=15, bold=True, color=COL_KI)


# ===========================================================================
# Slide 6: Generalisation (Edward + Haili)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'cs')
add_title(s, 'Same model, fresh batch of 30 cells: pipeline and findings both hold up')

# Left side
left_x, left_w = Inches(0.5), Inches(6.0)
add_textbox(s, left_x, Inches(1.5), left_w, Inches(2.0),
            'Mid-hackathon a second upload arrived: 30 fresh cells, 5,862 frames, same 6 categories. The fine-tuned model has never seen these cells. We re-ran the entire pipeline (segmentation, centering, migration, lamellipodia) with no retraining and no parameter changes.',
            font_size=14, color=COL_BODY)

gen_data = [
    ['Metric', 'Batch 1', 'Batch 2'],
    ['Empty-mask rate', '0.13%', '1.28%'],
    ['KI migration speed', '0.68 ± 0.08', '0.72 ± 0.10'],
    ['KI lam/cyto mean', '1.827', '1.830'],
    ["Cohen's d KI vs WT (mig)", '−1.40', '−1.22'],
]
add_table(s, left_x, Inches(3.4), left_w, Inches(2.0), gen_data,
          font_size=12, col_widths=[Inches(2.6), Inches(1.7), Inches(1.7)])

add_textbox(s, left_x, Inches(5.6), left_w, Inches(0.5),
            'Both findings reproduce on independent cells.',
            font_size=14, bold=True, color=COL_KI)
add_textbox(s, left_x, Inches(6.05), left_w, Inches(1.0),
            'KI lam/cyto means agree within 0.003 across batches. Pipeline + fine-tuned weights generalise; data-processing decisions (adaptive normalize, mask-stability filter, percentile=50 split) carry over without tuning.',
            font_size=11, color=COL_META)

# Right side: image
img_path = IMG_DIR / 'migration_trajectories_2nd.png'
if img_path.exists():
    add_image(s, img_path, Inches(6.9), Inches(1.5), height=Inches(5.0))
add_textbox(s, Inches(6.9), Inches(6.6), Inches(6.2), Inches(0.6),
            'Batch 2 trajectories. KI cells (last two panels) again cover noticeably less ground than WT or KO. Same pattern, fresh cells.',
            font_size=10, color=COL_META, align=PP_ALIGN.CENTER)


# ===========================================================================
# Slide 7: Tools (Edward)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'edward')
add_title(s, 'Drop-in pieces for future projects')

tools = [
    ('Adaptive normalize',
     'per-stack p99 detector decides whether to stretch. Fixes low-batch failure without breaking high-batch images. 100% concordance with the dataset\'s 1/2 folder labels.'),
    ('Self-training pipeline',
     'filter v3 outputs by mask area + IoU, retrain cyto3 on the surviving pairs. No hand-drawn labels required.'),
    ('Mask-stability filter',
     'rejects per-pair frames where IoU < 0.3 or area drops below 30% of per-cell median. Caught the wt1/33 fragmentation artefact.'),
    ('Generic compare_masks.py',
     'IoU, Dice, boundary IoU, Hausdorff 95, centroid distance. Drop in any reference masks (hand-drawn or otherwise).'),
    ('Generic plot_trajectories.py',
     '2×3 panel figure from any trajectories.json. Same script ran on batch 1 and batch 2.'),
]
y = Inches(1.5)
for name, desc in tools:
    add_textbox(s, Inches(0.5), y, Inches(3.3), Inches(0.8),
                name, font_size=15, bold=True, color=COL_EDWARD)
    add_textbox(s, Inches(3.9), y, Inches(9.0), Inches(0.85),
                desc, font_size=13, color=COL_BODY)
    y += Inches(0.92)

add_textbox(s, Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.5),
            'Reproducible: all code, manifests, processed CSV outputs, and the 26 MB fine-tuned weights are in the public repo.',
            font_size=14, bold=True, color=COL_DARK)
add_textbox(s, Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.4),
            'github.com/Haili321/bioimage-hackathon-beh',
            font_size=13, color=COL_META, align=PP_ALIGN.CENTER)


# ===========================================================================
# Slide 8: Caveats + future (all)
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'all')
add_title(s, 'Honest caveats + next steps')

add_textbox(s, Inches(0.5), Inches(1.45), Inches(6.0), Inches(0.4),
            'Caveats',
            font_size=18, bold=True, color=COL_DARK)
caveats = [
    ('Sample size',
     'N=10 cells per condition per batch. Effect sizes are large but a formal mixed-effects ANOVA with batch as a random effect is still pending.'),
    ('No hand-drawn ground truth',
     'the fine-tuned model learned the v3 consensus shape, not a real reference. compare_masks.py is ready to swap in real annotations.'),
    ('One stubborn case',
     'ko1/7 still failed at 28.3% empty (30 of 106 frames) on the fresh batch. Worth a manual look: bad recording, or genuine phenotype?'),
    ('Photobleaching',
     'not corrected. Absolute intensities should be read with caution. Ratios (lam/cyto) are partially robust because both populations decay together.'),
]
y = Inches(1.85)
for label, text in caveats:
    add_textbox(s, Inches(0.5), y, Inches(2.3), Inches(0.4),
                label, font_size=13, bold=True, color=COL_KI)
    add_textbox(s, Inches(2.85), y, Inches(4.0), Inches(1.3),
                text, font_size=11, color=COL_BODY)
    y += Inches(1.05)

add_textbox(s, Inches(7.3), Inches(1.45), Inches(6.0), Inches(0.4),
            'Next steps',
            font_size=18, bold=True, color=COL_DARK)
nexts = [
    'Mixed-effects model with batch as random effect (deliver p-values, not just effect sizes).',
    'Photobleaching correction for absolute intensity.',
    'Three-layer split (body / transition / lamellipodia) if biology team wants finer granularity.',
    'napari interactive viewer for the biology team to browse cells frame-by-frame.',
]
y = Inches(1.85)
for n in nexts:
    add_textbox(s, Inches(7.3), y, Inches(0.3), Inches(0.4),
                '•', font_size=14, bold=True, color=COL_BADEER)
    add_textbox(s, Inches(7.55), y, Inches(5.6), Inches(1.05),
                n, font_size=12, color=COL_BODY)
    y += Inches(1.05)


# ===========================================================================
# Slide 9: Closing + Q&A
# ===========================================================================
s = prs.slides.add_slide(BLANK)
add_speaker_badge(s, 'all')

add_textbox(s, Inches(0.5), Inches(1.5), SW - Inches(1), Inches(1.2),
            'Thank you',
            font_size=60, bold=True, color=COL_DARK, align=PP_ALIGN.CENTER)

add_textbox(s, Inches(0.5), Inches(3.3), SW - Inches(1), Inches(0.5),
            'Briefing:  www.dcs.warwick.ac.uk/~u1898019/hackathon-beh/',
            font_size=18, color=COL_BODY, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.5), Inches(3.9), SW - Inches(1), Inches(0.5),
            'GitHub:  github.com/Haili321/bioimage-hackathon-beh',
            font_size=18, color=COL_BODY, align=PP_ALIGN.CENTER)

add_textbox(s, Inches(0.5), Inches(4.85), SW - Inches(1), Inches(0.5),
            '17+ commits  ·  3 contributors  ·  fine-tuned weights and slides shipped',
            font_size=14, color=COL_META, align=PP_ALIGN.CENTER)

add_textbox(s, Inches(0.5), Inches(5.8), SW - Inches(1), Inches(1.0),
            'Questions?',
            font_size=44, bold=True, color=COL_KI, align=PP_ALIGN.CENTER)


# ===========================================================================
prs.save(OUT)
print(f'[done] saved {OUT}')
print(f'  slides: {len(prs.slides)}')
print(f'  size: {OUT.stat().st_size / 1024:.0f} KB')
